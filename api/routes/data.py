"""GET /data/comparison and /data/source-docs — vendor data endpoints for the UI."""

import base64
import json
import os
from pathlib import Path

from fastapi import APIRouter, Query
from src.db.connection import get_db
from src.ingestion.detector import detect_format

router = APIRouter()


@router.get("/comparison")
async def get_comparison(rfx_id: str = Query("RFX-001")):
    """Return all comparison rows for an RFx, grouped by vendor."""
    async with get_db() as db:
        cursor = await db.execute(
            """
            SELECT vendor_id, line_id, description, price_inr, unit_raw,
                   currency_raw, confidence, flags, extraction_status
            FROM comparison WHERE rfx_id = ?
            ORDER BY vendor_id, line_id
            """,
            (rfx_id,),
        )
        rows = await cursor.fetchall()

    grouped: dict[str, list] = {}
    for r in rows:
        vid = r["vendor_id"]
        if vid not in grouped:
            grouped[vid] = []
        grouped[vid].append({
            "line_id":          r["line_id"],
            "description":      r["description"],
            "price_inr":        round(r["price_inr"], 2) if r["price_inr"] is not None else None,
            "unit_raw":         r["unit_raw"],
            "currency_raw":     r["currency_raw"],
            "confidence":       round(r["confidence"], 2) if r["confidence"] is not None else None,
            "flags":            json.loads(r["flags"]) if r["flags"] else [],
            "extraction_status": r["extraction_status"],
        })

    return {"rfx_id": rfx_id, "vendors": grouped}


def _preview_text(file_path: str, fmt: str, max_chars: int = 600) -> str:
    """Return a raw text snippet from the vendor file."""
    try:
        if fmt == "text":
            with open(file_path, encoding="utf-8", errors="replace") as f:
                return f.read(max_chars)

        if fmt == "excel":
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            lines = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 8:
                    break
                cells = [str(c) if c is not None else "" for c in row]
                lines.append(" | ".join(cells))
            return "\n".join(lines)

        if fmt == "docx":
            from docx import Document
            doc = Document(file_path)
            parts = []
            for para in doc.paragraphs[:12]:
                if para.text.strip():
                    parts.append(para.text.strip())
            return "\n".join(parts)[:max_chars]

        if fmt == "pdf":
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                text = ""
                for page in pdf.pages[:2]:
                    text += (page.extract_text() or "") + "\n"
            return text[:max_chars]

    except Exception as exc:
        return f"[preview unavailable: {exc}]"

    return ""


def _preview_image_b64(file_path: str) -> str:
    """Return base64-encoded image for inline display."""
    try:
        with open(file_path, "rb") as f:
            data = f.read()
        return base64.standard_b64encode(data).decode("utf-8")
    except Exception:
        return ""


@router.get("/source-docs")
async def get_source_docs(vendor_dir: str = "data/vendor_responses"):
    """Return raw source file metadata and previews for all vendor files."""
    base = Path(vendor_dir)
    if not base.exists():
        return {"files": []}

    files = sorted(
        p for p in base.iterdir()
        if p.is_file() and "vendor_" in p.name and "response" in p.name
    )

    result = []
    for fp in files:
        fmt = detect_format(str(fp))
        size_kb = round(fp.stat().st_size / 1024, 1)
        stem = fp.stem
        vendor_id = stem
        for suffix in ("_response", "_quote", "_reply"):
            if suffix in stem:
                vendor_id = stem[: stem.index(suffix)]
                break

        entry = {
            "filename": fp.name,
            "vendor_id": vendor_id,
            "format": fmt,
            "extension": fp.suffix.lstrip(".").upper(),
            "size_kb": size_kb,
        }

        if fmt == "image":
            entry["preview_type"] = "image"
            entry["preview_b64"] = _preview_image_b64(str(fp))
            ext = fp.suffix.lstrip(".").lower()
            entry["mime"] = f"image/{ext}" if ext != "jpg" else "image/jpeg"
        else:
            entry["preview_type"] = "text"
            entry["preview_text"] = _preview_text(str(fp), fmt)

        result.append(entry)

    return {"files": result}


# ── Chart data endpoints ────────────────────────────────────────

@router.get("/chart/price-comparison")
async def chart_price_comparison(rfx_id: str = Query("RFX-001"), line_id: int = Query(1)):
    """Bar chart data: all vendors' price for one line item."""
    async with get_db() as db:
        cursor = await db.execute(
            """
            SELECT vendor_id, price_inr, confidence, flags
            FROM comparison
            WHERE rfx_id = ? AND line_id = ?
            ORDER BY price_inr ASC NULLS LAST
            """,
            (rfx_id, line_id),
        )
        rows = await cursor.fetchall()

        # Get description for this line
        c2 = await db.execute(
            "SELECT description FROM comparison WHERE rfx_id = ? AND line_id = ? LIMIT 1",
            (rfx_id, line_id),
        )
        desc_row = await c2.fetchone()

    return {
        "line_id": line_id,
        "description": desc_row["description"] if desc_row else f"Line {line_id}",
        "vendors": [
            {
                "vendor_id": r["vendor_id"],
                "price_inr": round(r["price_inr"], 2) if r["price_inr"] else None,
                "confidence": round(r["confidence"], 2) if r["confidence"] else 0,
                "flags": json.loads(r["flags"]) if r["flags"] else [],
            }
            for r in rows
        ],
    }


@router.get("/chart/vendor-scorecard")
async def chart_vendor_scorecard(rfx_id: str = Query("RFX-001")):
    """Scorecard: avg confidence + flag count + lines quoted per vendor."""
    async with get_db() as db:
        cursor = await db.execute(
            """
            SELECT vendor_id,
                   AVG(confidence) as avg_conf,
                   COUNT(*) as total_lines,
                   SUM(CASE WHEN price_inr IS NOT NULL THEN 1 ELSE 0 END) as priced_lines,
                   flags
            FROM comparison
            WHERE rfx_id = ?
            GROUP BY vendor_id
            ORDER BY avg_conf DESC
            """,
            (rfx_id,),
        )
        rows = await cursor.fetchall()

        # Count distinct flags per vendor
        c2 = await db.execute(
            "SELECT vendor_id, flags FROM comparison WHERE rfx_id = ?", (rfx_id,)
        )
        flag_rows = await c2.fetchall()

    flag_counts: dict[str, set] = {}
    for r in flag_rows:
        vid = r["vendor_id"]
        flags = json.loads(r["flags"]) if r["flags"] else []
        flag_counts.setdefault(vid, set()).update(flags)

    return {
        "vendors": [
            {
                "vendor_id": r["vendor_id"],
                "avg_confidence": round(r["avg_conf"] or 0, 2),
                "total_lines": r["total_lines"],
                "priced_lines": r["priced_lines"],
                "unique_flags": len(flag_counts.get(r["vendor_id"], set())),
                "flag_names": sorted(flag_counts.get(r["vendor_id"], set())),
            }
            for r in rows
        ]
    }


@router.get("/chart/coverage")
async def chart_coverage(rfx_id: str = Query("RFX-001")):
    """Coverage heatmap: for each vendor × line_id, whether price exists."""
    async with get_db() as db:
        cursor = await db.execute(
            """
            SELECT vendor_id, line_id, description, price_inr, confidence
            FROM comparison
            WHERE rfx_id = ?
            ORDER BY vendor_id, line_id
            """,
            (rfx_id,),
        )
        rows = await cursor.fetchall()

    vendors: list[str] = []
    lines: dict[int, str] = {}
    cells: dict[str, dict[int, dict]] = {}

    for r in rows:
        vid = r["vendor_id"]
        lid = r["line_id"]
        if vid not in vendors:
            vendors.append(vid)
        lines[lid] = r["description"] or f"Line {lid}"
        cells.setdefault(vid, {})[lid] = {
            "has_price": r["price_inr"] is not None,
            "confidence": round(r["confidence"] or 0, 2),
        }

    return {
        "vendors": vendors,
        "lines": [{"id": lid, "description": desc[:40]} for lid, desc in sorted(lines.items())],
        "cells": cells,
    }
